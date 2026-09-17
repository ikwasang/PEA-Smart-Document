"""Local Excel imports: validate first, then atomically replace one data group."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from datetime import datetime, timezone
import hashlib,json,re,threading,time,uuid
import openpyxl

KINDS={'budgets':'Data ของงบแฟ้มงาน','accounts':'รหัสบัญชี','people':'รายชื่อบุคลากร'}
LOCK=threading.Lock()
PENDING={}
def val(v):
 if isinstance(v,(int,float)) and not isinstance(v,bool):return str(int(v)) if v==int(v) else str(v)
 return str(v or '').strip()
def counts(data):
 return {'budgetFiles':len(data.get('budgetFiles',[])),'networks':sum(len(f['networks']) for f in data.get('budgetFiles',[])),'accounts':len(data.get('accounts',[])),'costCenters':len(data.get('costCenters',[])),'people':len(data.get('people',[]))}

def parse_excel(kind,raw):
 if kind not in KINDS:raise ValueError('ประเภทข้อมูลไม่ถูกต้อง')
 try:
  with ZipFile(BytesIO(raw)) as z:
   if sum(i.file_size for i in z.infolist())>80*1024*1024:raise ValueError('ไฟล์ Excel มีขนาดใหญ่เกินไป')
  book=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True)
 except (BadZipFile,KeyError,OSError) as e:raise ValueError('อ่านไฟล์ไม่ได้ กรุณาใช้ Excel .xlsx ตามต้นแบบ') from e
 warnings=[];records={};centers={};network_owners={};people=[]
 try:
  for sheet in book:
   if (sheet.max_row or 0)>30000 or (sheet.max_column or 0)>200:raise ValueError('ชีตมีขนาดใหญ่เกินขอบเขตข้อมูลต้นแบบ')
   if kind=='budgets':
    cells=list(sheet.iter_rows());rows=[[c.value for c in row] for row in cells]
    for i,row in enumerate(rows):
     code=val(row[3]) if len(row)>3 else ''
     if not code.startswith(('I-','C-')):continue
     if i+4>=len(rows):raise ValueError(f'ข้อมูลแฟ้มงานไม่ครบในชีต {sheet.title} แถว {i+1}')
     network=val(rows[i+2][3]) if len(rows[i+2])>3 else ''
     expenses=[val(v) for v in rows[i+4][1:5] if v]
     if not re.fullmatch(r'\d{10}',network) or not expenses:raise ValueError(f'โครงข่ายหรือประเภทค่าไม่ครบในชีต {sheet.title} แถว {i+1}')
     if network in network_owners and network_owners[network]!=code:raise ValueError(f'โครงข่าย {network} ซ้ำข้ามแฟ้มงาน')
     network_owners[network]=code
     entry=records.setdefault(code,{'code':code,'name':val(row[0]),'sheet':sheet.title,'networks':[],'expenses':[]})
     new={'code':network,'name':val(rows[i+2][0]),'expenses':expenses,'workName':val(row[0]),'approval':val(rows[i+1][0])}
     activity_budgets={}
     for budget_index in range(i+5,min(i+10,len(rows))):
      if val(rows[budget_index][3] if len(rows[budget_index])>3 else '').startswith(('I-','C-')):break
      if not val(rows[budget_index][0]).startswith('งบประมาณ'):continue
      color=cells[budget_index][0].fill.fgColor
      rgb=color.rgb[-6:] if color.type=='rgb' else ''
      if len(rgb)!=6:continue
      red,green,blue=[int(rgb[k:k+2],16) for k in (0,2,4)]
      if not (green>red+25 and green>blue+25):continue
      for column in range(1,min(5,len(rows[i+4]))):
       activity=val(rows[i+4][column]);amount=rows[budget_index][column]
       if activity and isinstance(amount,(int,float)) and not isinstance(amount,bool):activity_budgets[activity]=str(amount)
      new['budgetSourceRow']=budget_index+1
      break
     new['activityBudgets']=activity_budgets
     used={}
     end=next((j for j in range(i+1,len(rows)) if len(rows[j])>3 and val(rows[j][3]).startswith(('I-','C-'))),len(rows))
     for j in range(i+5,end):
      if val(rows[j][0])=='ยอดเงินที่ใช้ไป':
       for column in range(1,min(5,len(rows[i+4]),len(rows[j]))):
        activity=val(rows[i+4][column]);amount=rows[j][column]
        if activity and isinstance(amount,(int,float)) and not isinstance(amount,bool):used[activity]=str(amount)
       new['spentSourceRow']=j+1
       break
     new['activitySpent']=used
     old=next((n for n in entry['networks'] if n['code']==network),None)
     if old and old!=new:
      old['conflict']=True;old['activityBudgets']={};old['activitySpent']={}
      warnings.append(f'โครงข่าย {network} ข้อมูลซ้ำไม่ตรงกันในชีต {sheet.title} กรุณาตรวจต้นทางก่อนใช้ยอดเงิน')
     if not old:entry['networks'].append(new)
     entry['expenses']=list(dict.fromkeys(entry['expenses']+expenses))
   elif kind=='accounts':
    for row in sheet.values:
     if len(row)>2 and re.fullmatch(r'\d{8}',val(row[1])):
      code,name=val(row[1]),val(row[2])
      if not name:raise ValueError(f'รหัสบัญชี {code} ไม่มีชื่อบัญชี')
      if code in records and records[code]['name']!=name:raise ValueError(f'ชื่อบัญชี {code} ไม่ตรงกันระหว่างชีต')
      records[code]={'code':code,'name':name}
     for v in row:
      if isinstance(v,str):
       for code in re.findall(r'A\d{9}',v):centers[code]={'code':code,'name':v.replace('\n',' ').strip()}
   else:
    header=False
    for row in sheet.values:
     if len(row)<3:continue
     if 'ชื่อ' in val(row[1]) and 'ตำแหน่ง' in val(row[2]):header=True;continue
     if not header or not val(row[1]):continue
     name,position=val(row[1]),val(row[2])
     if not position:raise ValueError(f'บุคลากร {name} ไม่มีตำแหน่ง')
     if len(name)>40:raise ValueError(f'ชื่อ {name} ยาวเกิน 40 ตัวอักษรที่แบบฟอร์มรองรับ')
     if any(p['name']==name for p in people):
      warnings.append(f'พบชื่อ {name} ซ้ำในชีต {sheet.title} จะเก็บทุกรายการตามไฟล์ รวมตำแหน่งของแต่ละรายการ')
     people.append({'name':name,'position':position})
 finally:book.close()
 if not (people if kind=='people' else records):raise ValueError('ไม่พบข้อมูลตามประเภทที่เลือก กรุณาตรวจไฟล์และรูปแบบคอลัมน์ต้นแบบ')
 key={'budgets':'budgetFiles','accounts':'accounts','people':'people'}[kind]
 result={key:people if kind=='people' else list(records.values())}
 if kind=='accounts':
  if centers:result['costCenters']=list(centers.values())
  else:warnings.append('ไม่พบรหัสศูนย์ต้นทุนในไฟล์นี้ จะคงศูนย์ต้นทุนเดิมไว้')
 return result,warnings

def preview(root,kind,raw,filename):
 changes,warnings=parse_excel(kind,raw)
 with LOCK:
  path=root/'master-data.json';original=path.read_bytes();before=json.loads(original);after={**before,**changes}
  for token in list(PENDING):
   if PENDING[token]['expires']<time.time():del PENDING[token]
  if len(PENDING)>=8:PENDING.pop(next(iter(PENDING)))
  token=uuid.uuid4().hex
  PENDING[token]={'changes':changes,'kind':kind,'filename':filename[:200],'fingerprint':hashlib.sha256(original).hexdigest(),'expires':time.time()+900}
 return {'token':token,'before':counts(before),'after':counts(after),'warnings':warnings,'sample':next(iter(changes.values()))[:5],'kind':kind,'filename':filename}

def commit(root,token):
 with LOCK:
  pending=PENDING.get(token)
  if not pending or pending['expires']<time.time():raise ValueError('ผลตรวจไฟล์หมดอายุ กรุณาตรวจไฟล์อีกครั้ง')
  path=root/'master-data.json';original=path.read_bytes()
  if hashlib.sha256(original).hexdigest()!=pending['fingerprint']:raise ValueError('ข้อมูลมีการเปลี่ยนแปลงแล้ว กรุณาตรวจไฟล์อีกครั้งก่อนอัปเดต')
  data=json.loads(original);data.update(pending['changes']);now=datetime.now(timezone.utc).isoformat()
  data.setdefault('updates',{})[pending['kind']]={'filename':pending['filename'],'updatedAt':now}
  backup=root/'local-data'/'master-backups';backup.mkdir(parents=True,exist_ok=True)
  backupfile=backup/(datetime.now().strftime('%Y%m%d-%H%M%S')+'-'+uuid.uuid4().hex[:8]+'.json');backupfile.write_bytes(original)
  temporary=path.with_name('master-data.'+uuid.uuid4().hex+'.tmp')
  try:temporary.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8');temporary.replace(path)
  finally:
   if temporary.exists():temporary.unlink()
  del PENDING[token]
  return {'counts':counts(data),'backup':backupfile.name,'updatedAt':now}
