from pathlib import Path
import openpyxl,json,re
root=Path('D:/SuperAng project/ไฟล์ต้นแบบสำหรับ export')
def value(v):return str(int(v)) if isinstance(v,(int,float)) else str(v or '').strip()
accounts={};centers={};people={};files=[]
w=openpyxl.load_workbook(root/'รหัสบัญชี.xlsx',data_only=True,read_only=True)
for sheet in w:
 for row in sheet.values:
  if len(row)>2 and re.fullmatch(r'\d{8}',value(row[1])) and row[2]:accounts[value(row[1])]=value(row[2])
  for v in row:
   if isinstance(v,str):
    for code in re.findall(r'A\d{9}',v):centers[code]=v.replace('\n',' ').strip()
w=openpyxl.load_workbook(root/'รายชื่อ ประธานกรรมการ, กรรมการ และเจ้าหน้าที่จัดทำเอกสาร.xlsx',data_only=True,read_only=True)
for sheet in w:
 for row in sheet.iter_rows(min_row=2,values_only=True):
  if row[1]:people[value(row[1])]={'name':value(row[1]),'position':value(row[2])}
w=openpyxl.load_workbook(root/'Data ของงบแฟ้มงาน.xlsx',data_only=True,read_only=True)
for sheet in w:
 rows=list(sheet.values)
 for index,row in enumerate(rows):
  code=value(row[3]) if len(row)>3 else ''
  if not code.startswith('I-') or index+4>=len(rows):continue
  network=value(rows[index+2][3])
  if not re.fullmatch(r'\d{10}',network):continue
  expenses=[value(v) for v in rows[index+4][1:5] if v]
  entry=next((f for f in files if f['code']==code),None)
  if entry is None:
   entry={'code':code,'name':value(row[0]),'sheet':sheet.title,'networks':[],'expenses':[]};files.append(entry)
  if not any(n['code']==network for n in entry['networks']):entry['networks'].append({'code':network,'name':value(rows[index+2][0]),'expenses':expenses})
  entry['expenses']=list(dict.fromkeys(entry['expenses']+expenses))
data={'accounts':[{'code':k,'name':v} for k,v in accounts.items()],'costCenters':[{'code':k,'name':v} for k,v in centers.items()],'people':list(people.values()),'budgetFiles':files,'sources':[p.name for p in root.glob('*.xlsx')]}
Path('prototype/master-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:len(data[k]) for k in ['accounts','costCenters','people','budgetFiles']}))
